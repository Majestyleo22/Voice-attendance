"""
Export Helpers for Voice Attendance System
Generates professional CSV, Excel (.xlsx), and PDF reports.
"""

import io
import csv
from datetime import datetime


def generate_csv(records: list) -> io.BytesIO:
    """Generates a CSV file buffer from attendance records."""
    output = io.StringIO()
    writer = csv.writer(output)

    # Header
    writer.writerow([
        "S.No",
        "Register Number",
        "Student Name",
        "Department",
        "Section",
        "Year",
        "Date",
        "Time",
        "Status",
        "Recognized Speech",
    ])

    for i, r in enumerate(records, 1):
        writer.writerow([
            i,
            r.get("register_number", ""),
            r.get("name", ""),
            r.get("department", ""),
            r.get("section", ""),
            r.get("year", ""),
            r.get("date", ""),
            r.get("time", ""),
            r.get("status", "Present"),
            r.get("recognized_text", ""),
        ])

    mem = io.BytesIO()
    mem.write(output.getvalue().encode("utf-8-sig"))
    mem.seek(0)
    return mem


def generate_excel(records: list, filter_summary: str = "") -> io.BytesIO:
    """Generates a formatted Excel (.xlsx) workbook using openpyxl."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    # Title styling
    title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Title Block
    ws.merge_cells("A1:J1")
    ws["A1"] = "VOICE ATTENDANCE SYSTEM - ATTENDANCE REPORT"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:J2")
    ws["A2"] = f"Generated on {datetime.now().strftime('%Y-%m-%d %I:%M %p')} | Total Records: {len(records)} {filter_summary}"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Headers
    headers = [
        "S.No", "Register No", "Student Name", "Department",
        "Section", "Year", "Date", "Time", "Status", "Recognized Utterance"
    ]
    ws.append([]) # Blank row 3
    ws.append(headers) # Row 4
    ws.row_dimensions[4].height = 25

    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data Rows
    for idx, r in enumerate(records, 1):
        row_data = [
            idx,
            r.get("register_number", ""),
            r.get("name", ""),
            r.get("department", ""),
            r.get("section", ""),
            r.get("year", ""),
            r.get("date", ""),
            r.get("time", ""),
            r.get("status", "Present"),
            r.get("recognized_text", "")
        ]
        ws.append(row_data)
        current_row = idx + 4
        ws.row_dimensions[current_row].height = 20
        for col_num in range(1, len(row_data) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", horizontal="center" if col_num in (1, 2, 4, 5, 6, 7, 8, 9) else "left")
            if idx % 2 == 0:
                cell.fill = alt_fill

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if cell.row > 2: # Skip merged title rows
                max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    mem = io.BytesIO()
    wb.save(mem)
    mem.seek(0)
    return mem


def generate_pdf(records: list, filter_summary: str = "") -> io.BytesIO:
    """Generates a professional PDF report using ReportLab."""
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    mem = io.BytesIO()
    doc = SimpleDocTemplate(
        mem,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1E3A8A"),
        alignment=1, # Center
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        "DocSub",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748B"),
        alignment=1,
        spaceAfter=15
    )
    cell_style = ParagraphStyle(
        "Cell",
        parent=styles["Normal"],
        fontSize=9,
        leading=11,
        alignment=0
    )
    header_cell_style = ParagraphStyle(
        "HeaderCell",
        parent=styles["Normal"],
        fontSize=9,
        leading=11,
        textColor=colors.white,
        fontName="Helvetica-Bold",
        alignment=1
    )

    elements = []
    elements.append(Paragraph("Voice Attendance System – Official Attendance Report", title_style))
    elements.append(Paragraph(
        f"Generated on {datetime.now().strftime('%B %d, %Y at %I:%M %p')} | Total Verified Entries: {len(records)} {filter_summary}",
        subtitle_style
    ))

    # Table data
    headers = [
        Paragraph("<b>#</b>", header_cell_style),
        Paragraph("<b>Reg No</b>", header_cell_style),
        Paragraph("<b>Student Name</b>", header_cell_style),
        Paragraph("<b>Dept</b>", header_cell_style),
        Paragraph("<b>Sec</b>", header_cell_style),
        Paragraph("<b>Year</b>", header_cell_style),
        Paragraph("<b>Date</b>", header_cell_style),
        Paragraph("<b>Time</b>", header_cell_style),
        Paragraph("<b>Status</b>", header_cell_style),
        Paragraph("<b>Voice Input</b>", header_cell_style),
    ]

    table_data = [headers]

    for idx, r in enumerate(records, 1):
        row = [
            Paragraph(str(idx), cell_style),
            Paragraph(str(r.get("register_number", "")), cell_style),
            Paragraph(str(r.get("name", "")), cell_style),
            Paragraph(str(r.get("department", "")), cell_style),
            Paragraph(str(r.get("section", "")), cell_style),
            Paragraph(str(r.get("year", "")), cell_style),
            Paragraph(str(r.get("date", "")), cell_style),
            Paragraph(str(r.get("time", "")), cell_style),
            Paragraph(f"<font color='green'><b>{r.get('status', 'Present')}</b></font>", cell_style),
            Paragraph(str(r.get("recognized_text", ""))[:40], cell_style),
        ]
        table_data.append(row)

    col_widths = [30, 60, 140, 50, 40, 45, 75, 75, 65, 160]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))

    elements.append(t)
    doc.build(elements)
    mem.seek(0)
    return mem
